"""export.py — convert extracted samples to Excel."""

import io
from itertools import chain
from typing import Any, Dict, List


def samples_to_excel(samples: List[Dict[str, Any]]) -> bytes:
    """Convert a list of extracted samples to an .xlsx file (raw bytes)."""
    try:
        import openpyxl
        from openpyxl.utils import get_column_letter
    except ImportError:
        raise ImportError("openpyxl is required for Excel export: pip install openpyxl")

    # -----------------------------------------------------------------------
    # STEP 1 — scan all samples for max matrices, max fillers, all prop keys
    # -----------------------------------------------------------------------
    if not samples:
        max_matrices = 0
        max_fillers = 0
        all_prop_keys: List[str] = []
    else:
        max_matrices = max(
            len(s.get("matrix_name") or []) for s in samples
        )
        max_fillers = max(
            len(_real_fillers(s)) for s in samples
        )
        prop_keys = set()
        for s in samples:
            for k in (s.get("properties") or {}).keys():
                prop_keys.add(k)
        all_prop_keys = sorted(prop_keys)

    # -----------------------------------------------------------------------
    # STEP 2 — build the column list
    # -----------------------------------------------------------------------
    columns: List[str] = ["source_pdf", "sample_id"]
    for i in range(1, max_matrices + 1):
        columns.extend([
            f"matrix_{i}_name",
            f"matrix_{i}_amount",
            f"matrix_{i}_unit",
            f"matrix_{i}_description",
            f"matrix_{i}_surface_treatment",
        ])
    for i in range(1, max_fillers + 1):
        columns.extend([
            f"filler_{i}_name",
            f"filler_{i}_amount",
            f"filler_{i}_unit",
            f"filler_{i}_description",
            f"filler_{i}_surface_treatment",
        ])
    columns.extend(["thickness_value", "thickness_unit"])
    for p in all_prop_keys:
        columns.extend([f"{p}_value", f"{p}_unit", f"{p}_conditions"])

    # -----------------------------------------------------------------------
    # STEP 3 — emit one row per property measurement
    # -----------------------------------------------------------------------
    rows: List[List[Any]] = []
    for sample in samples:
        fixed = _fixed_cells(sample, max_matrices, max_fillers)
        props = sample.get("properties") or {}
        if not props:
            row = [fixed.get(c, "") for c in columns]
            rows.append(row)
            continue

        for p_key, p_data in props.items():
            row_dict = dict(fixed)
            if isinstance(p_data, dict):
                row_dict[f"{p_key}_value"] = p_data.get("value", "")
                row_dict[f"{p_key}_unit"] = p_data.get("unit", "")
                row_dict[f"{p_key}_conditions"] = _flatten_conditions(
                    p_data.get("conditions") or []
                )
            row = [row_dict.get(c, "") for c in columns]
            rows.append(row)

    # -----------------------------------------------------------------------
    # STEP 4 — write workbook with auto-fit column widths + frozen header
    # -----------------------------------------------------------------------
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Extracted Data"

    for row_idx, row in enumerate(chain([columns], rows), start=1):
        for col_idx, value in enumerate(row, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            if isinstance(value, str):
                # Article/model text (including property-derived headers) is
                # data, never an Excel formula. Preserve the literal value.
                cell.data_type = "s"

    MAX_WIDTH = 60
    PADDING = 2
    for col_idx, col_name in enumerate(columns, start=1):
        max_len = len(str(col_name))
        for row in rows:
            v = row[col_idx - 1]
            if v is None or v == "":
                continue
            max_len = max(max_len, len(str(v)))
        width = min(max_len, MAX_WIDTH) + PADDING
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    ws.freeze_panes = "A2"

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.read()


# ---------------------------------------------------------------------------
# helpers
# ---------------------------------------------------------------------------
def _real_fillers(sample: Dict[str, Any]) -> List[Dict[str, Any]]:
    """Filter out placeholder filler entries (none / not specified)."""
    return [
        f for f in (sample.get("fillers") or [])
        if f.get("name", "").lower() not in ("none", "not specified", "")
    ]


def _fixed_cells(
    sample: Dict[str, Any],
    max_matrices: int,
    max_fillers: int,
) -> Dict[str, Any]:
    """Return the fixed-column cells for one sample (sample_id, matrices,
    fillers, thickness). These values are identical across every row that
    belongs to the same sample."""
    out: Dict[str, Any] = {
        "source_pdf": sample.get("_source_pdf", ""),
        "sample_id": sample.get("sample_id", ""),
    }

    matrices = sample.get("matrix_name") or []
    for i in range(max_matrices):
        m = matrices[i] if i < len(matrices) else {}
        comp = m.get("composition") or {}
        out[f"matrix_{i+1}_name"] = m.get("name", "")
        out[f"matrix_{i+1}_amount"] = comp.get("amount", "")
        out[f"matrix_{i+1}_unit"] = comp.get("type", "")
        out[f"matrix_{i+1}_description"] = m.get("description", "") or ""
        out[f"matrix_{i+1}_surface_treatment"] = m.get("surface_treatment", "") or ""

    fillers = _real_fillers(sample)
    for i in range(max_fillers):
        f = fillers[i] if i < len(fillers) else {}
        comp = f.get("composition") or {}
        out[f"filler_{i+1}_name"] = f.get("name", "")
        out[f"filler_{i+1}_amount"] = comp.get("amount", "")
        out[f"filler_{i+1}_unit"] = comp.get("type", "")
        out[f"filler_{i+1}_description"] = f.get("description", "") or ""
        out[f"filler_{i+1}_surface_treatment"] = f.get("surface_treatment", "") or ""

    thick = sample.get("sample_thickness") or {}
    out["thickness_value"] = thick.get("value", "")
    out["thickness_unit"] = thick.get("unit", "")
    return out


def _flatten_conditions(conds: List[Dict[str, Any]]) -> str:
    """Render a list of condition dicts as 'type=value unit; type=value unit'."""
    if not conds:
        return ""
    parts = []
    for c in conds:
        t = c.get("type", "") or ""
        v = c.get("value", "")
        u = c.get("unit", "") or ""
        parts.append(f"{t}={v} {u}".strip())
    return "; ".join(parts)
